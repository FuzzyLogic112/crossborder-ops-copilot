# -*- coding: utf-8 -*-
"""生成 1688 询价记录表（带公式，边问边算到仓成本）。

为什么要带公式：跟店家聊的时候就能看到「这家到底是不是最便宜」，
而不是回来再跑一遍面板。公式和 src/sourcing.py 的口径完全一致：

    实际采购量 = max(你的采购量, MOQ)          ← MOQ 高于需求就得多买
    适用单价   = 按实际采购量落在哪个阶梯
    含税单价   = 需要专票且报价不含票 ? 单价×(1+开票点数) : 单价
    运费/件    = 含运 ? 0 : 整批运费 ÷ 实际采购量
    一次性/件  = (打样费 + 模具费) ÷ 实际采购量
    到仓成本   = 含税单价 + 运费/件 + 一次性/件

用法：
    python make_quote_sheet.py                      # 生成到 reports/
    python make_quote_sheet.py --out D:/某路径.xlsx
    python make_quote_sheet.py --product "保温饭盒包" --qty 300
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).parent

HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
CALC_FILL = PatternFill("solid", fgColor="EAF3F6")
PARAM_FILL = PatternFill("solid", fgColor="FFF4D6")
WARN_FONT = Font(color="B00020", bold=True)
NOTE_FONT = Font(color="666666", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (标题, 列宽, 是否为计算列, 提示)
COLUMNS = [
    ("店铺名",        22, False, "能认出是哪一家就行"),
    ("平台",          10, False, "1688 / 义乌购 / 阿里国际站 / 线下档口"),
    ("标价¥",         9,  False, "商品页挂的单价"),
    ("200件价¥",      10, False, "阶梯价。没有阶梯就三格都填标价"),
    ("500件价¥",      10, False, ""),
    ("1000件价¥",     11, False, ""),
    ("MOQ",           8,  False, "最小起订量。「10件起批」就填 10"),
    ("含票",          7,  False, "报价是否已含增值税。1688 多数不含"),
    ("开票加点%",     11, False, "不含票时要加几个点，通常 6~13"),
    ("含运",          7,  False, "是否包邮到你手上"),
    ("运费¥",         10, False, "整批总额，不是每件！"),
    ("单个净重(克)",  13, False, "跨境按公斤计费，这个数很关键"),
    ("打样费¥",       10, False, "一次性，会摊到每件"),
    ("模具费¥",       10, False, "定制才有，现货一般 0"),
    ("交期(天)",      10, False, "超过 20 天注意断货"),
    ("报价日期",      12, False, "报价会变，没日期不能当决策依据"),
    ("链接",          26, False, ""),
    ("备注",          20, False, "材质、克重、是否工厂等"),
    ("实际采购量",    11, True,  "MOQ 高于你的采购量时要多买"),
    ("适用单价¥",     11, True,  "按实际采购量落在哪个阶梯"),
    ("含税单价¥",     11, True,  ""),
    ("运费/件¥",      11, True,  ""),
    ("一次性/件¥",    12, True,  "打样费+模具费摊销"),
    ("到仓成本¥",     12, True,  "★ 真正该比的就是这一列"),
    ("提醒",          40, True,  "自动检查，别忽略"),
]

HEADER_ROW = 7          # 表头所在行
FIRST_DATA = 8          # 第一条数据行
N_ROWS = 8              # 预留几行


def col(name):
    return get_column_letter([c[0] for c in COLUMNS].index(name) + 1)


def build(product, qty, need_invoice, vat_pct):
    wb = Workbook()
    ws = wb.active
    ws.title = "询价记录"

    # ── 参数区 ──
    ws["A1"] = "1688 询价记录表"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "黄色格子是参数，改了下面的到仓成本会自动重算。蓝色格子是公式，不要手动改。"
    ws["A2"].font = NOTE_FONT

    params = [
        ("A4", "产品", "B4", product,
         "C4", "你要问的是什么品"),
        ("A5", "你的采购量(件)", "B5", qty,
         "C5", "先小批量测款，200~300 是常见起步量"),
        ("D4", "需要专票", "E4", "是" if need_invoice else "否",
         "F4", "出口退税通常需要"),
        ("D5", "开票税率%", "E5", vat_pct,
         "F5", "店家没说加几个点时按这个算"),
    ]
    for lk, lv, vk, vv, nk, nv in params:
        ws[lk] = lv
        ws[lk].font = Font(bold=True, size=10)
        ws[vk] = vv
        ws[vk].fill = PARAM_FILL
        ws[vk].border = BORDER
        ws[nk] = nv
        ws[nk].font = NOTE_FONT

    dv_yn_param = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
    ws.add_data_validation(dv_yn_param)
    dv_yn_param.add(ws["E4"])

    # ── 表头 ──
    for i, (title, width, is_calc, hint) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=title)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
        if hint:
            c.comment = None  # 提示放到「怎么填」页，批注在手机上看不到
    ws.row_dimensions[HEADER_ROW].height = 30

    # ── 数据行与公式 ──
    P_QTY, P_INV, P_VAT = "$B$5", "$E$4", "$E$5"
    c_price, c_t200, c_t500, c_t1000 = col("标价¥"), col("200件价¥"), col("500件价¥"), col("1000件价¥")
    c_moq, c_tax, c_pt = col("MOQ"), col("含票"), col("开票加点%")
    c_fin, c_frt, c_wt = col("含运"), col("运费¥"), col("单个净重(克)")
    c_smp, c_mold, c_lead = col("打样费¥"), col("模具费¥"), col("交期(天)")
    c_date, c_shop = col("报价日期"), col("店铺名")
    c_aq, c_up, c_tp = col("实际采购量"), col("适用单价¥"), col("含税单价¥")
    c_fpu, c_opu, c_land = col("运费/件¥"), col("一次性/件¥"), col("到仓成本¥")
    c_warn = col("提醒")

    dv_yn = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv_yn)

    for r in range(FIRST_DATA, FIRST_DATA + N_ROWS):
        blank = '{c}{r}=""'.format(c=c_shop, r=r)

        # 实际采购量 = max(采购量, MOQ)
        ws["%s%d" % (c_aq, r)] = ('=IF({b},"",MAX({q},IF({m}{r}="",1,{m}{r})))'
                                  .format(b=blank, q=P_QTY, m=c_moq, r=r))
        # 适用单价：按实际采购量落在哪个阶梯，阶梯空着就回落到标价
        ws["%s%d" % (c_up, r)] = (
            '=IF({b},"",'
            'IF(AND({aq}{r}>=1000,{t3}{r}<>""),{t3}{r},'
            'IF(AND({aq}{r}>=500,{t2}{r}<>""),{t2}{r},'
            'IF({t1}{r}<>"",{t1}{r},{p}{r}))))'
            .format(b=blank, aq=c_aq, t1=c_t200, t2=c_t500, t3=c_t1000, p=c_price, r=r))
        # 含税单价：需要专票且报价不含票 → 加点（店家给了加点用店家的，否则用默认税率）
        ws["%s%d" % (c_tp, r)] = (
            '=IF({b},"",IF(AND({inv}="是",{tax}{r}="否"),'
            '{up}{r}*(1+IF({pt}{r}="",{vat},{pt}{r})/100),{up}{r}))'
            .format(b=blank, inv=P_INV, tax=c_tax, up=c_up, pt=c_pt, vat=P_VAT, r=r))
        # 运费/件
        ws["%s%d" % (c_fpu, r)] = (
            '=IF({b},"",IF({fin}{r}="是",0,IF({aq}{r}=0,0,IF({frt}{r}="",0,{frt}{r})/{aq}{r})))'
            .format(b=blank, fin=c_fin, frt=c_frt, aq=c_aq, r=r))
        # 一次性摊销/件
        ws["%s%d" % (c_opu, r)] = (
            '=IF({b},"",IF({aq}{r}=0,0,(IF({s}{r}="",0,{s}{r})+IF({m}{r}="",0,{m}{r}))/{aq}{r}))'
            .format(b=blank, aq=c_aq, s=c_smp, m=c_mold, r=r))
        # 到仓成本
        ws["%s%d" % (c_land, r)] = (
            '=IF({b},"",ROUND({tp}{r}+{fpu}{r}+{opu}{r},2))'
            .format(b=blank, tp=c_tp, fpu=c_fpu, opu=c_opu, r=r))
        # 提醒：把 src/sourcing.py 里那几条警告搬过来
        ws["%s%d" % (c_warn, r)] = (
            '=IF({b},"",TRIM('
            'IF({m}{r}>{q},"⚠MOQ高于采购量，多出的是压货风险 ","")&'
            'IF(AND({inv}="是",{tax}{r}="否"),"⚠不含票已加税 ","")&'
            'IF(AND({fin}{r}="否",OR({frt}{r}="",{frt}{r}=0)),"⚠标不含运但运费填0，成本被低估 ","")&'
            'IF({d}{r}="","⚠没有报价日期 ","")&'
            'IF({w}{r}="","⚠没问净重，算不了国际运费 ","")&'
            'IF(AND({l}{r}<>"",{l}{r}>20),"⚠交期偏长 ","")'
            '))'
            .format(b=blank, m=c_moq, q=P_QTY, inv=P_INV, tax=c_tax,
                    fin=c_fin, frt=c_frt, d=c_date, w=c_wt, l=c_lead, r=r))

        for i in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=i)
            cell.border = BORDER
            if COLUMNS[i - 1][2]:
                cell.fill = CALC_FILL
        ws["%s%d" % (c_warn, r)].font = WARN_FONT
        ws["%s%d" % (c_warn, r)].alignment = Alignment(wrap_text=True, vertical="top")
        ws["%s%d" % (c_land, r)].font = Font(bold=True)
        for c in (c_tax, c_fin):
            dv_yn.add(ws["%s%d" % (c, r)])

    last = FIRST_DATA + N_ROWS - 1
    rng = "%s%d:%s%d" % (c_land, FIRST_DATA, c_land, last)
    # 最低到仓成本标绿 —— 这才是该选的那家，不是标价最低的
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['AND({c}{r}<>"",{c}{r}=MIN({rng}))'.format(c=c_land, r=FIRST_DATA, rng="$%s$%d:$%s$%d" % (c_land, FIRST_DATA, c_land, last))],
        fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(bold=True, color="006100")))
    # 标价最低的标黄，用来对照「标价陷阱」
    rng_p = "$%s$%d:$%s$%d" % (c_price, FIRST_DATA, c_price, last)
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (c_price, FIRST_DATA, c_price, last),
        FormulaRule(formula=['AND({c}{r}<>"",{c}{r}=MIN({rng}))'.format(c=c_price, r=FIRST_DATA, rng=rng_p)],
                    fill=PatternFill("solid", fgColor="FFEB9C")))

    # ── 结论区 ──
    cr = last + 2
    ws["A%d" % cr] = "结论"
    ws["A%d" % cr].font = Font(bold=True, size=12)
    ws["A%d" % (cr + 1)] = "标价最低的是"
    ws["B%d" % (cr + 1)] = ('=IFERROR(INDEX({s}{f}:{s}{l},MATCH(MIN({p}{f}:{p}{l}),{p}{f}:{p}{l},0)),"")'
                            .format(s=c_shop, p=c_price, f=FIRST_DATA, l=last))
    ws["A%d" % (cr + 2)] = "到仓最低的是"
    ws["B%d" % (cr + 2)] = ('=IFERROR(INDEX({s}{f}:{s}{l},MATCH(MIN({d}{f}:{d}{l}),{d}{f}:{d}{l},0)),"")'
                            .format(s=c_shop, d=c_land, f=FIRST_DATA, l=last))
    ws["A%d" % (cr + 3)] = "是否踩了标价陷阱"
    ws["B%d" % (cr + 3)] = ('=IF(OR(B{a}="",B{b}=""),"",IF(B{a}=B{b},"没有，标价最低的确实到仓也最低",'
                            '"⚠ 踩了！标价最低的不是到仓最低的，按标价选会选错"))'
                            .format(a=cr + 1, b=cr + 2))
    for i in range(1, 4):
        ws["A%d" % (cr + i)].font = Font(bold=True, size=10)
        ws["B%d" % (cr + i)].fill = CALC_FILL
        ws["B%d" % (cr + i)].border = BORDER
    ws["B%d" % (cr + 3)].font = WARN_FONT
    ws.merge_cells("B%d:F%d" % (cr + 3, cr + 3))

    ws.freeze_panes = "C%d" % FIRST_DATA

    # ── 说明页 ──
    ws2 = wb.create_sheet("怎么填")
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 62
    ws2.column_dimensions["C"].width = 26
    rows = [("列", "填什么", "不知道就填")]
    rows += [(t, h, "") for t, _, is_calc, h in COLUMNS if not is_calc and h]
    defaults = {"MOQ": "1", "含票": "否（算高一点更安全）", "含运": "否",
                "运费¥": "0，但记得回头补", "打样费¥": "0", "模具费¥": "0",
                "报价日期": "今天", "开票加点%": "留空，按上面的税率算"}
    for i, (a, b, _) in enumerate(rows):
        c = defaults.get(a, "")
        ws2.append([a, b, c] if i else [a, b, "不知道就填"])
    for cell in ws2[1]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=3):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws2.append([])
    tips = [
        ("最容易漏的两项", ""),
        ("含不含票", "1688 多数标价不含税。同样标价 ¥28：含票到仓 ¥28；"
                     "不含票且你要专票 → ¥31.64。13% 足以翻转三家排名。"
                     "直接问一句「这个价含不含票」。"),
        ("单个净重", "别接受「挺轻的」。追一句「麻烦称一下单个净重，多少克就行，"
                     "我要算国际运费」。软货还要问外箱尺寸和装箱数，容易抛重。"),
        ("", ""),
        ("运费填总额", "店家说「运费 300」就填 300，别自己先除以件数——"
                       "表格会按实际采购量摊。"),
        ("供应商填店名", "填「平邮」「顺丰」是不对的，那属于运费信息，"
                         "该体现在「含运」和「运费」两列。"),
        ("至少问三家", "只有一家比不出标价陷阱。实测有组报价里标价 ¥23.80 的"
                       "到仓 ¥30.35，反而高于标价 ¥29.00 的那家。"),
        ("", ""),
        ("这张表算的是什么", "到仓成本 = 含税单价 + 运费/件 + (打样费+模具费)/件。"
                             "口径和项目里 src/sourcing.py 完全一致，"
                             "填完可以直接抄进面板「④ 供应商比价」核对。"),
        ("算不了什么", "国际头程运费。那要等你拿到净重，再去面板「⑥ 定价计算器」算。"),
    ]
    for a, b in tips:
        ws2.append([a, b])
        if a and not b:
            ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True, size=11)
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    return wb


def main():
    ap = argparse.ArgumentParser(description="生成 1688 询价记录表")
    ap.add_argument("--product", default="保温饭盒包（保温午餐袋）")
    ap.add_argument("--qty", type=int, default=300)
    ap.add_argument("--no-invoice", action="store_true", help="不需要增值税专用发票")
    ap.add_argument("--vat", type=float, default=13, help="开票税率%%，默认 13")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    out = Path(a.out) if a.out else BASE / "reports" / "询价记录表.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = build(a.product, a.qty, not a.no_invoice, a.vat)
    wb.save(out)
    print("已生成：%s" % out)
    print("  产品：%s   采购量：%d 件   需要专票：%s"
          % (a.product, a.qty, "否" if a.no_invoice else "是"))
    print("  预留 %d 行，填满了直接复制一行往下拉。" % N_ROWS)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
